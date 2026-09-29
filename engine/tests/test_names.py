"""표시 이름 적재 (D-075) — 합성 데이터. 실제 코드네임·제품군 이름은 테스트에 쓰지 않는다."""
import openpyxl
from scm_engine import names as N


def _key(tmp_path):
    wb = openpyxl.Workbook(); wb.active.title = "안내"
    c = wb.create_sheet("기종 코드네임"); c.append(["원본 코드네임", "변환 코드"]); c.append(["ALPHA", "MDL901"]); c.append(["BETA", "MDL902"])
    f = wb.create_sheet("Family 전체값"); f.append(["원본 값", "변환 값", "출처 파일"])
    for row in [("#N/A", "#N/A", "x"), ("Alpha Mono(A100)", "MDL901 Mono(A237)", "MC_OL_vs_ACT"), ("ALPHA MONO(A100)", "MDL901 MONO(A237)", "부품_XCN"),
                ("10 SERIES", "47 SERIES", "부품"), ("SAME", "SAME", "부품"), ("Beta's 2", "MDL902-2", "옵션"), ("Beta 2 old", "MDL902-2", "옵션"),
                ("GAMMA PRO", "MDL903 PRO", "부품_XCN"), ("Gamma Pro", "MDL903 PRO", "옵션, MC_OL_vs_ACT"), (None, None, None)]:
        f.append(list(row))
    wb.create_sheet("부품·옵션 코드").append(["원본 코드", "변환 코드"])
    p = tmp_path / "key.xlsx"; wb.save(p); return p


def test_read_pairs_takes_names_only(tmp_path):
    df, stat = N.read_pairs(_key(tmp_path))
    fam = df[df.kind == "family"]
    assert set(df.kind) == {"family", "codename"} and len(df[df.kind == "codename"]) == 2
    assert "#N/A" not in set(fam.anon) and "SAME" not in set(fam.anon)                       # 빈 값·같은 값은 넣지 않는다
    assert dict(zip(fam.anon, fam.real_name))["MDL901 Mono(A237)"] == "Alpha Mono(A100)"      # 대소문자가 다른 쌍은 각각 보존
    assert dict(zip(fam.anon, fam.real_name))["MDL901 MONO(A237)"] == "ALPHA MONO(A100)"
    assert list(fam[fam.anon == "MDL902-2"].real_name) == ["Beta's 2"] and stat["duplicate_anon"] == 2   # 같은 익명 이름은 첫 쌍
    assert list(fam[fam.anon == "MDL903 PRO"].real_name) == ["Gamma Pro"]                       # 단, 기종 OL 파일에 나온 표기를 우선
    assert dict(zip(df.anon, df.anon_key))["MDL901 Mono(A237)"] == "MDL901 MONO(A237)"


def test_to_sql_replaces_whole_table_and_escapes(tmp_path):
    df, _ = N.read_pairs(_key(tmp_path))
    sql = N.to_sql(df)
    assert sql.startswith("begin;") and "delete from app.name_alias;" in sql and sql.rstrip().endswith("commit;")
    assert "('family','MDL902-2','MDL902-2','Beta''s 2')" in sql
    assert "부품·옵션 코드" not in sql and "0000-" not in sql                                  # 코드 복원 쌍은 다루지 않는다
