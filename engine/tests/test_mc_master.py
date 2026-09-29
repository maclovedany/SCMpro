"""기종(MC) 마스터·OL 실적 적재 (D-077) — 합성 데이터. 실제 제품군 이름·코드는 테스트에 쓰지 않는다."""
import openpyxl
import pandas as pd
from scm_engine import mc_master as M


def _book(tmp_path):
    wb = openpyxl.Workbook(); wb.active.title = "Summary"
    ln = wb.create_sheet("전임 후속기")
    ln.append([None] * 6); ln.append([None, "MC Category", "before Family", "보안용 Family 변경", "After Family", "보안용 Family 변경"])
    ln.append([None, "DT", "Alpha One(100)", "AL1", "Beta Two(200)", "BT2"])
    ln.append([None, "DT", None, None, "Beta Mono", "BTM"])
    ln.append([None, "PRT", "Gamma SFP", "GMS", None, None])     # 회사 정리본은 이름에서 (MOQ n) 을 뗐다
    a = wb.create_sheet("FY23")
    a.append([None] * 13); a.append([None, None, None, "23' 4月", None, None, None, None, "23' 5月"])
    a.append(["Item Code", None, "Family", "Sales OL (1st)", "SCM OL", "Apr Act", "Act vs Sales OL", "Act vs SCM OL", "Sales OL (1st)", "SCM OL", "May Act", "Act vs Sales OL", "Act vs SCM OL", "Sales OL (1st)", "SCM OL", "FY Act"])     # 끝의 세 열은 연간 합계
    a.append(["TL000001", "DT", "Alpha One(100)", 10, 20, 15, 1.5, 0.75, 5, 6, None, None, None, 15, 26, 15])
    a.append(["TL000009", "PRT", "Gamma SFP  (MOQ 24)", 0, 0, 0, "-", "-", 1, 2, 3, 3, 1.5, 1, 2, 3])
    a.append([None, None, "DT/GC SUB TOTAL", 10, 20, 15, None, None, 6, 8, 3])
    a.append([0, None, "GRAND Total", 10, 20, 15, None, None, 6, 8, 3])
    b = wb.create_sheet("FY26-to202606")
    b.append([None] * 8); b.append([None, None, None, "26' 4月", None, None, None, "26' 5月"])
    b.append(["Item Code", None, "Family", "Sales\nOL", "SCM\nOL", "Apr\nAct", "Sales\nAct rate (%)", "Sales\nOL", "SCM\nOL", "May\nAct", "Sales\nAct rate (%)"])     # 실제 파일처럼 칸 안 줄바꿈
    b.append(["TL000002", "DT", "Beta Two(200)", 22, 110, 39, 1.77, 37, 47, 29, 0.78])
    b.append(["TL000003", "DT", "Beta Mono", 1, 2, 3, 3, None, None, None, None])
    b.append(["TX999999", "GC", "Delta 4C", 4, 5, 6, 1.5, 0, 0, 0, 0])
    c = wb.create_sheet("FY25")     # 구분이 넷째 열에 있는 시트
    c.append([None] * 9); c.append([None, None, None, None, "25' 4月"])
    c.append(["Item Code", None, "Family", "BIZ", "Sales OL", "SCM OL", "Apr Act", "Act vs Sales OL", "Act vs SCM OL"])
    c.append(["TL000002", None, "Beta Two(200)", "DT", 7, 8, 9, 1.2, 1.1])
    p = tmp_path / "mc.xlsx"; wb.save(p); return p


KEY = M.Key(code={"TL000001": "TL900001", "TL000002": "TL900002", "TL000003": "TL900003", "TL000009": "TL900009"},
            family={"ALPHA ONE(100)": "MDL901 One(237)", "BETA TWO(200)": "MDL902 Two(337)", "GAMMA SFP (MOQ 24)": "MDL903 SFP (MOQ 161)", "DELTA 4C": "MDL904 4C"},
            codename={"ALPHA": "MDL901", "BETA": "MDL902", "GAMMA": "MDL903", "DELTA": "MDL904"})
MODELS = pd.DataFrame([("MDL901 One(237)", "MDL901", "TL900001"), ("MDL902 Two(337)", "MDL902", "TL900002"), ("MDL903 SFP (MOQ 161)", "MDL903", None)], columns=["model_key", "model_base", "iot_code"])


def test_read_plan_parses_month_blocks_and_skips_totals(tmp_path):
    df = M.read_plan(_book(tmp_path))
    assert set(df.family) == {"Alpha One(100)", "Gamma SFP  (MOQ 24)", "Beta Two(200)", "Beta Mono", "Delta 4C"}          # 소계·합계 행 제외
    r = df[(df.family == "Alpha One(100)")].set_index("ym")
    assert list(r.index) == ["2023-04", "2023-05"] and r.loc["2023-04", ["sales_ol", "scm_ol", "act"]].tolist() == [10, 20, 15]
    assert pd.isna(r.loc["2023-05", "act"]) and r.loc["2023-05", ["sales_ol", "scm_ol"]].tolist() == [5, 6]       # 빈 칸은 값 없음, 뒤의 연간 합계 열은 읽지 않음
    assert df[(df.family == "Beta Two(200)") & (df.ym == "2025-04")].iloc[0][["biz", "sales_ol", "act"]].tolist() == ["DT", 7, 9]   # 구분이 넷째 열
    assert len(df[df.family == "Beta Mono"]) == 1                                                                 # 값이 전부 빈 달은 만들지 않는다


def test_build_uses_company_alias_and_keeps_codes_anonymous(tmp_path):
    p = _book(tmp_path)
    fam, items, rep = M.build(M.read_plan(p), M.read_lineage(p), KEY, MODELS)
    f = fam.set_index("family_key")
    assert set(f.index) == {"AL1", "BT2", "BTM", "GMS", "Delta 4C"}                                              # 회사 약자, 약자 없는 것은 회사 파일의 이름 그대로 (D-078)
    assert f.loc["BT2", ["biz", "item_code", "model_key", "model_base", "predecessor", "has_alias"]].tolist() == ["DT", "TL900002", "MDL902 Two(337)", "MDL902", "AL1", True]
    assert f.loc["BTM", "model_base"] == "MDL902" and pd.isna(f.loc["BTM", "predecessor"])                       # 이름 앞 코드네임으로 기종 묶음
    assert f.loc["Delta 4C", ["biz", "has_alias", "model_base"]].tolist() == ["GC", False, "MDL904"] and pd.isna(f.loc["Delta 4C", "item_code"])   # 치환표에 없는 코드는 넣지 않는다
    assert fam.family_key.tolist() == ["BT2", "BTM", "Delta 4C", "AL1", "GMS"]                                  # 가장 최근 시트의 행 순서, 옛 시트에만 있는 것은 뒤
    assert f.loc["GMS", ["biz", "model_key", "model_base"]].tolist() == ["PRT", "MDL903 SFP (MOQ 161)", "MDL903"]      # (MOQ n) 이 붙은 옛 이름과 연결
    assert items[(items.family_key == "BT2")].sort_values("ym").ym.tolist() == ["2025-04", "2026-04", "2026-05"]
    assert rep["no_alias"] == ["Delta 4C"] and rep["unmapped_codes"] == 1
    blob = fam.to_csv() + items.to_csv()
    for real in ["Alpha", "Beta", "Gamma", "TL000001", "TL000002", "TX999999"]: assert real not in blob   # 약자가 있는 Family 의 실제 이름과 실코드는 결과에 없다


def test_to_sql_is_rerunnable(tmp_path):
    p = _book(tmp_path)
    fam, items, _ = M.build(M.read_plan(p), M.read_lineage(p), KEY, MODELS)
    sql = M.to_sql(fam, items)
    assert sql.startswith("begin;") and "delete from app.mc_plan_item;" in sql and "delete from app.mc_family;" in sql and sql.rstrip().endswith("commit;")
    assert "('BT2','DT','TL900002','MDL902 Two(337)','MDL902','AL1',true," in sql and "('BT2','2026-04',22,110,39)" in sql
    assert "('BTM','2026-04',1,2,3)" in sql and "null" in sql
