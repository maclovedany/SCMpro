"""기종(MC) 마스터·OL 실적 적재 (D-077, R-FC-16). 회사가 보낸 정리본(실제 이름·실코드)을 **로컬에서** 읽어
- 제품군 이름은 회사가 준 보안용 약자('전임 후속기' 시트)로,
- Item Code 는 치환표의 익명 코드로
바꾼 뒤 app.mc_family · app.mc_plan_item 에 넣는다. 실제 이름과 실코드는 DB·저장소에 남기지 않는다 (D-059, D-075)."""
from __future__ import annotations
import re
from dataclasses import dataclass
from pathlib import Path
import pandas as pd

LINEAGE_SHEET = "전임 후속기"
TOTAL_RE = re.compile(r"SUB\s*TOTAL|GRAND\s*TOTAL", re.I)
MONTH_RE = re.compile(r"(\d{2})\s*['’]\s*(\d{1,2})\s*月")
BIZ = ("DT", "GC", "PRT")
BLOCK_MAX = 5                 # 한 달 묶음의 최대 열 수 (Sales OL · SCM OL · Act · 비율 2)
MOQ_RE = re.compile(r"\s*\(MOQ[^)]*\)\s*$", re.I)


def _s(v) -> str:
    return "" if v is None else str(v).strip()

def _norm(v) -> str:
    return re.sub(r"\s+", " ", _s(v)).upper()

def _base(v) -> str:
    """회사 정리본은 이름 끝의 (MOQ n) 을 뗐다 — 옛 이름과 맞추는 열쇠"""
    return MOQ_RE.sub("", _norm(v))

def _hdr(v) -> str:
    return re.sub(r"\s+", " ", _s(v))

def _num(v):
    return float(v) if isinstance(v, (int, float)) and not isinstance(v, bool) else None


@dataclass
class Key:
    """치환표: 실코드 → 익명 코드, 실제 제품군 이름 → 익명 이름, 코드네임 → 기종 묶음(MDLnnn)"""
    code: dict[str, str]
    family: dict[str, str]
    codename: dict[str, str]

    @classmethod
    def load(cls, path: Path) -> "Key":
        import openpyxl
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        def pairs(sheet: str, norm):
            return {norm(r[0]): _s(r[1]) for r in list(wb[sheet].iter_rows(values_only=True))[1:] if r and len(r) >= 2 and _s(r[0]) and _s(r[1]) and _s(r[0]) != "#N/A"}
        return cls(code=pairs("부품·옵션 코드", _s), family=pairs("Family 전체값", _norm), codename=pairs("기종 코드네임", _norm))

    def model_base(self, family: str) -> str | None:
        """이름 앞의 코드네임(긴 것 우선) → 기종 묶음"""
        toks = re.split(r"[\s(/_-]+", _norm(family))
        for n in range(min(3, len(toks)), 0, -1):
            for t in (" ".join(toks[:n]), re.sub(r"\d+$", "", " ".join(toks[:n]))):
                if t and t in self.codename: return self.codename[t]
        return None


def _sheets(path: Path) -> dict[str, list[list]]:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    return {ws.title: [list(r) for r in ws.iter_rows(values_only=True)] for ws in wb.worksheets}


def read_plan(path: Path) -> pd.DataFrame:
    """FY 시트 → (fy_sheet, item_code, biz, family, ym, sales_ol, scm_ol, act). 월 묶음은 머리글 위 줄의 "23' 4月" 표기로 찾는다 — 합계 열은 읽지 않는다."""
    out = []
    for name, rows in _sheets(path).items():
        if not name.upper().startswith("FY"): continue
        hi = next((i for i, r in enumerate(rows) if r and _hdr(r[0]) == "Item Code"), None)
        if hi is None or hi == 0: continue
        hdr = [_hdr(c) for c in rows[hi]]; months = rows[hi - 1]     # 머리글은 칸 안에서 줄이 바뀌어 있다 ("Sales\nOL")
        fam_col = hdr.index("Family"); biz_col = hdr.index("BIZ") if "BIZ" in hdr else 1
        blocks = []                                                   # (ym, sales_col, scm_col, act_col)
        for j, c in enumerate(months):
            m = MONTH_RE.search(_s(c))
            if not m: continue
            ym = f"20{m.group(1)}-{int(m.group(2)):02d}"
            nxt = next((k for k in range(j + 1, len(months)) if MONTH_RE.search(_s(months[k]))), len(hdr))
            cols: dict[str, int] = {}
            for k in range(j, min(nxt, j + BLOCK_MAX, len(hdr))):          # 묶음 안에서 처음 나오는 열만 — 마지막 달 뒤의 연간 합계 열을 읽지 않는다
                h = hdr[k]; kind = "sales" if h.startswith("Sales OL") else "scm" if h.startswith("SCM OL") else "act" if h.endswith(" Act") else None
                if kind and kind not in cols: cols[kind] = k
            if {"sales", "scm", "act"} <= set(cols): blocks.append((ym, cols["sales"], cols["scm"], cols["act"]))
        for r in rows[hi + 1:]:
            fam = _s(r[fam_col]) if len(r) > fam_col else ""
            if not fam or TOTAL_RE.search(fam): continue
            code = _s(r[0]); biz = _s(r[biz_col]) if len(r) > biz_col else ""
            for ym, a, b, c in blocks:
                v = [_num(r[k]) if len(r) > k else None for k in (a, b, c)]
                if all(x is None for x in v): continue
                out.append((name, code if code and code != "0" else None, biz if biz in BIZ else None, fam, ym, *v))
    return pd.DataFrame(out, columns=["fy_sheet", "item_code", "biz", "family", "ym", "sales_ol", "scm_ol", "act"])


def read_lineage(path: Path) -> pd.DataFrame:
    """'전임 후속기' 시트 → (biz, before, before_alias, after, after_alias)"""
    rows = _sheets(path)[LINEAGE_SHEET]
    hi = next(i for i, r in enumerate(rows) if any(_s(c) == "MC Category" for c in r))
    c0 = [_s(c) for c in rows[hi]].index("MC Category")
    out = [(_s(r[c0]) or None, _s(r[c0 + 1]) or None, _s(r[c0 + 2]) or None, _s(r[c0 + 3]) or None, _s(r[c0 + 4]) or None)
           for r in rows[hi + 1:] if r and len(r) > c0 + 4 and (_s(r[c0 + 1]) or _s(r[c0 + 3]))]
    return pd.DataFrame(out, columns=["biz", "before", "before_alias", "after", "after_alias"])


def build(plan: pd.DataFrame, lineage: pd.DataFrame, key: Key, models: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    """plan·lineage(실제 이름·실코드) → families(family_key …) · items(family_key, ym, …). 결과에는 약자와 익명 코드만 남는다."""
    alias, biz_of, pred = {}, {}, {}
    for r in lineage.itertuples():
        for real, al in ((r.before, r.before_alias), (r.after, r.after_alias)):
            if real and al: alias[_base(real)] = al; biz_of[al] = r.biz if r.biz in BIZ else None
        if r.before_alias and r.after_alias: pred[r.after_alias] = r.before_alias
    by_iot = {m.iot_code: m for m in models.itertuples() if isinstance(m.iot_code, str) and m.iot_code}
    by_key = {_norm(m.model_key): m for m in models.itertuples()}
    old_name = {}                                                 # (MOQ n) 을 뗀 실제 이름 → 익명 이름
    for real, anon in key.family.items(): old_name.setdefault(_base(real), anon)
    fams, no_alias, unmapped = {}, [], set()
    last = plan.sort_values(["fy_sheet", "ym"]).groupby("family", sort=False).last().reset_index()      # 제품군마다 가장 최근 시트의 코드·구분
    order: dict[str, int] = {}                                     # 회사 파일의 행 순서: 가장 최근 시트부터, 그 시트에 없는 제품군은 뒤에
    for sh in sorted(plan.fy_sheet.unique(), reverse=True):
        for f in plan[plan.fy_sheet == sh].family.drop_duplicates(): order.setdefault(f, len(order))
    for r in last.itertuples():
        anon_name = key.family.get(_norm(r.family)) or old_name.get(_base(r.family))
        fk = alias.get(_base(r.family)) or anon_name or (f"MC-{key.code[r.item_code]}" if r.item_code in key.code else None)
        if not fk: continue
        has_alias = _base(r.family) in alias
        if not has_alias: no_alias.append(fk)
        iot = key.code.get(r.item_code) if r.item_code else None
        if r.item_code and not iot: unmapped.add(r.item_code)
        m = by_iot.get(iot) if iot else None
        if m is None and anon_name: m = by_key.get(_norm(anon_name))
        fams[r.family] = {"family_key": fk, "biz": biz_of.get(fk) or r.biz, "item_code": iot, "model_key": getattr(m, "model_key", None),
                          "model_base": (getattr(m, "model_base", None) or key.model_base(r.family)), "predecessor": pred.get(fk), "has_alias": has_alias, "sort_no": order.get(r.family, 0)}
    fam = pd.DataFrame(fams.values()).drop_duplicates("family_key").sort_values("sort_no").reset_index(drop=True)
    fam.loc[~fam.predecessor.isin(set(fam.family_key)), "predecessor"] = None                       # 실적 파일에 없는 전임기는 연결하지 않는다
    items = plan.assign(family_key=plan.family.map(lambda f: fams.get(f, {}).get("family_key"))).dropna(subset=["family_key"])
    items = items.sort_values(["fy_sheet", "ym"]).drop_duplicates(["family_key", "ym"], keep="last")[["family_key", "ym", "sales_ol", "scm_ol", "act"]].reset_index(drop=True)
    return fam, items, {"families": len(fam), "rows": len(items), "no_alias": sorted(no_alias), "unmapped_codes": len(unmapped)}


def _q(v) -> str:
    if v is None or (isinstance(v, float) and pd.isna(v)): return "null"
    if isinstance(v, bool): return "true" if v else "false"
    if isinstance(v, (int, float)): return str(int(v)) if float(v).is_integer() else repr(float(v))
    return "'" + str(v).replace("'", "''") + "'"


def to_sql(fam: pd.DataFrame, items: pd.DataFrame) -> str:
    """두 테이블 전체 교체 (재실행 안전)"""
    f = ",\n".join("(" + ",".join(_q(v) for v in (r.family_key, r.biz, r.item_code, r.model_key, r.model_base, r.predecessor, bool(r.has_alias), int(r.sort_no))) + ")" for r in fam.itertuples())
    rows = ["(" + ",".join(_q(v) for v in (r.family_key, r.ym, r.sales_ol, r.scm_ol, r.act)) + ")" for r in items.itertuples()]
    body = "".join("insert into app.mc_plan_item(family_key, ym, sales_ol, scm_ol, act) values\n" + ",\n".join(rows[i:i + 1000]) + ";\n" for i in range(0, len(rows), 1000))
    return ("begin;\ndelete from app.mc_plan_item;\ndelete from app.mc_family;\n"
            "insert into app.mc_family(family_key, biz, item_code, model_key, model_base, predecessor, has_alias, sort_no) values\n" + f + ";\n" + body + "commit;\n")
